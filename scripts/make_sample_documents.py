"""Create the sample documents in data/documents (a Word guide and a PDF sheet).

They are fictional NetCo / FibreLink documents used to demonstrate document processing: upload
them in the chat UI (or let scripts/e2e_check.py do it) and ask about their content.

    uv run --with python-docx --with fpdf2 python scripts/make_sample_documents.py
"""

from pathlib import Path

from docx import Document
from fpdf import FPDF

OUT = Path(__file__).resolve().parents[1] / "data" / "documents"


def make_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("FibreLink ONT-100 quick guide", level=0)
    doc.add_paragraph(
        "The FibreLink ONT-100 is the optical network terminal that NetCo installs at fiber "
        "customers on a GPON line. This guide is for the support team."
    )
    doc.add_heading("Factory reset", level=1)
    doc.add_paragraph(
        "To restore the ONT-100 to factory settings, hold the recessed reset button for "
        "12 seconds with a paperclip. The ONT restarts within three minutes and must then be "
        "provisioned again by the access team. Reset the ONT only when the access team asks "
        "for it."
    )
    doc.add_heading("Optical power levels", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Receive power", "Meaning"
    for level, meaning in [
        ("-8 to -27 dBm", "Normal receive power"),
        ("Below -27 dBm", "Signal too weak: check the fiber patch cable, then book an engineer"),
        ("Above -8 dBm", "Signal too strong (overload)"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text = level, meaning
    doc.add_heading("Handling fiber connectors", level=1)
    doc.add_paragraph(
        "Never look into the end of a fiber or an ONT port. Keep the dust caps on unused "
        "connectors and do not bend a patch cable tighter than 30 mm."
    )
    doc.save(path)


def make_pdf(path: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 12, "FibreLink ONT-200 installation sheet", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(
        0, 6,
        "The ONT-200 connects customers to an XGS-PON line. Install it near the fiber wall "
        "outlet and a power socket.",
        new_x="LMARGIN", new_y="NEXT",
    )
    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Power supply", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(
        0, 6, "Use only the supplied 12 V DC power adapter. The power light is green when the "
        "ONT has power.", new_x="LMARGIN", new_y="NEXT",
    )
    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "PON light", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=11)
    with pdf.table(col_widths=(40, 60)) as table:
        header = table.row()
        header.cell("PON light")
        header.cell("Meaning")
        for light, meaning in [
            ("Blinking green", "Ranging: the ONT is searching for the OLT"),
            ("Solid green", "Registered at the OLT"),
            ("Off", "No optical signal"),
        ]:
            row = table.row()
            row.cell(light)
            row.cell(meaning)
    pdf.output(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    make_docx(OUT / "FibreLink-ONT-100-quick-guide.docx")
    make_pdf(OUT / "FibreLink-ONT-200-installation-sheet.pdf")
    print(f"written to {OUT}")
